# ============================================================
# BAYER LEVERKUSEN - PROFESSIONAL EXCEL REPORT
# ============================================================

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
from openpyxl.worksheet.table import Table, TableStyleInfo
import os


# ------------------------------------------------------------
# FILE PATHS
# ------------------------------------------------------------

PLAYER_FILE = "data/player_metrics.csv"
DEPENDENCY_FILE = "data/squad_dependency.csv"

OUTPUT_FILE = "Bayer_Leverkusen_Squad_Risk_Analysis.xlsx"


# ------------------------------------------------------------
# LEVERKUSEN STYLE
# ------------------------------------------------------------

RED = "E32221"
DARK_RED = "B71918"
BLACK = "111111"
DARK_GRAY = "444444"
LIGHT_GRAY = "E7E7E7"
VERY_LIGHT_GRAY = "F5F5F5"
WHITE = "FFFFFF"

thin_gray = Side(
    style="thin",
    color="D9D9D9"
)


# ------------------------------------------------------------
# LOAD DATA
# ------------------------------------------------------------

print("\nLoading analysis data...")

players = pd.read_csv(PLAYER_FILE)
dependency = pd.read_csv(DEPENDENCY_FILE)

print("Player metrics loaded:", len(players))
print("Squad dependency loaded:", len(dependency))


# ------------------------------------------------------------
# CLEAN PLAYER NAMES
# ------------------------------------------------------------

name_map = {
    "Alejandro Grimaldo García": "Alejandro Grimaldo",
    "Exequiel Alejandro Palacios": "Exequiel Palacios",
    "Edmond Fayçal Tapsoba": "Edmond Tapsoba",
    "Piero Martín Hincapié Reyna": "Piero Hincapié",
    "Victor Okoh Boniface": "Victor Boniface",
    "Odilon Kossonou": "Odilon Kossounou",
    "Arthur Augusto de Matos Soares": "Arthur",
    "Gustavo Adolfo Puerta Molano": "Gustavo Puerta"
}

players["player"] = players["player"].replace(name_map)
dependency["player"] = dependency["player"].replace(name_map)


# ------------------------------------------------------------
# CREATE WORKBOOK
# ------------------------------------------------------------

wb = Workbook()

# Remove default sheet
default_sheet = wb.active
wb.remove(default_sheet)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def add_title(ws, title, subtitle):

    ws.merge_cells("A1:J1")
    ws["A1"] = title

    ws["A1"].font = Font(
        size=20,
        bold=True,
        color=WHITE
    )

    ws["A1"].fill = PatternFill(
        "solid",
        fgColor=RED
    )

    ws["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    ws["A2"] = subtitle

    ws["A2"].font = Font(
        size=11,
        italic=True,
        color=DARK_GRAY
    )

    ws["A2"].alignment = Alignment(
        horizontal="left"
    )

    ws.row_dimensions[2].height = 22


def style_header(cell):

    cell.fill = PatternFill(
        "solid",
        fgColor=BLACK
    )

    cell.font = Font(
        bold=True,
        color=WHITE
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    cell.border = Border(
        bottom=thin_gray
    )


def auto_width(ws, max_width=28):

    for column_cells in ws.columns:

        length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:

                if cell.value is not None:

                    length = max(
                        length,
                        len(str(cell.value))
                    )

            except Exception:
                pass

        ws.column_dimensions[column_letter].width = min(
            length + 3,
            max_width
        )


def add_dataframe(
    ws,
    dataframe,
    start_row,
    table_name
):

    # Headers
    for col_num, column_name in enumerate(
        dataframe.columns,
        start=1
    ):

        cell = ws.cell(
            row=start_row,
            column=col_num,
            value=column_name
        )

        style_header(cell)

    # Data
    for row_num, row in enumerate(
        dataframe.itertuples(index=False),
        start=start_row + 1
    ):

        for col_num, value in enumerate(
            row,
            start=1
        ):

            cell = ws.cell(
                row=row_num,
                column=col_num,
                value=value
            )

            cell.border = Border(
                bottom=thin_gray
            )

            if row_num % 2 == 0:

                cell.fill = PatternFill(
                    "solid",
                    fgColor=VERY_LIGHT_GRAY
                )

    end_row = start_row + len(dataframe)
    end_col = len(dataframe.columns)

    table_range = (
        f"A{start_row}:"
        f"{get_column_letter(end_col)}{end_row}"
    )

    table = Table(
        displayName=table_name,
        ref=table_range
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False
    )

    table.tableStyleInfo = style

    ws.add_table(table)


# ============================================================
# SHEET 1 - EXECUTIVE SUMMARY
# ============================================================

ws = wb.create_sheet(
    "Executive Summary"
)

ws.sheet_view.showGridLines = False

add_title(
    ws,
    "BAYER LEVERKUSEN | SQUAD RISK & RECRUITMENT ANALYSIS",
    "2023/24 Bundesliga • StatsBomb Open Data"
)


# Business problem

ws.merge_cells("A4:J4")

ws["A4"] = "BUSINESS PROBLEM"

ws["A4"].font = Font(
    size=12,
    bold=True,
    color=WHITE
)

ws["A4"].fill = PatternFill(
    "solid",
    fgColor=BLACK
)


ws.merge_cells("A5:J7")

ws["A5"] = (
    "Following Bayer Leverkusen's 2023/24 Bundesliga season, "
    "the recruitment department must identify areas where "
    "important team functions are concentrated among individual "
    "players. The analysis evaluates squad dependency and "
    "translates the highest attacking concentration into a "
    "measurable depth-recruitment profile."
)

ws["A5"].alignment = Alignment(
    wrap_text=True,
    vertical="top"
)

ws["A5"].font = Font(
    size=11,
    color=DARK_GRAY
)


# Key finding header

ws.merge_cells("A9:J9")

ws["A9"] = "KEY FINDINGS"

ws["A9"].font = Font(
    size=12,
    bold=True,
    color=WHITE
)

ws["A9"].fill = PatternFill(
    "solid",
    fgColor=BLACK
)


# Finding cards

findings = [
    (
        "Highest Attacking Dependency",
        "Florian Wirtz",
        "13.7%"
    ),
    (
        "Final-Third Carrying Share",
        "Florian Wirtz",
        "15.6%"
    ),
    (
        "Chance Creation Share",
        "Florian Wirtz",
        "14.7%"
    ),
    (
        "Final-Third Passing Share",
        "Florian Wirtz",
        "12.9%"
    )
]


start_row = 11

for title, player, value in findings:

    ws.merge_cells(
        start_row=start_row,
        start_column=1,
        end_row=start_row,
        end_column=3
    )

    title_cell = ws.cell(
        row=start_row,
        column=1
    )

    title_cell.value = title

    title_cell.font = Font(
        bold=True,
        color=WHITE
    )

    title_cell.fill = PatternFill(
        "solid",
        fgColor=RED
    )

    title_cell.alignment = Alignment(
        horizontal="center"
    )


    ws.merge_cells(
        start_row=start_row + 1,
        start_column=1,
        end_row=start_row + 1,
        end_column=2
    )

    player_cell = ws.cell(
        row=start_row + 1,
        column=1
    )

    player_cell.value = player

    player_cell.font = Font(
        size=12,
        bold=True
    )

    player_cell.alignment = Alignment(
        horizontal="center"
    )


    value_cell = ws.cell(
        row=start_row + 1,
        column=3
    )

    value_cell.value = value

    value_cell.font = Font(
        size=16,
        bold=True,
        color=RED
    )

    value_cell.alignment = Alignment(
        horizontal="center"
    )

    start_row += 3


# Recruitment recommendation

ws.merge_cells("A24:J24")

ws["A24"] = "RECRUITMENT PRIORITY"

ws["A24"].font = Font(
    size=12,
    bold=True,
    color=WHITE
)

ws["A24"].fill = PatternFill(
    "solid",
    fgColor=BLACK
)


ws.merge_cells("A25:J28")

ws["A25"] = (
    "The analysis indicates that attacking depth recruitment "
    "should prioritize players capable of progressing possession "
    "through carries, creating shots for teammates, delivering "
    "passes into the final third, and maintaining credible xG "
    "threat. The objective is not to find an identical replacement "
    "for Florian Wirtz, but to reduce concentration in the attacking "
    "functions most dependent on him."
)

ws["A25"].alignment = Alignment(
    wrap_text=True,
    vertical="top"
)

ws["A25"].font = Font(
    size=11,
    color=DARK_GRAY
)


for col in range(1, 11):

    ws.column_dimensions[
        get_column_letter(col)
    ].width = 14


# ============================================================
# SHEET 2 - PLAYER METRICS
# ============================================================

ws = wb.create_sheet(
    "Player Metrics"
)

ws.sheet_view.showGridLines = False

add_title(
    ws,
    "BAYER LEVERKUSEN | PLAYER METRICS",
    "Event-level player output • 2023/24 Bundesliga"
)


player_display = players[
    [
        "player",
        "passes",
        "carries",
        "shots",
        "xg",
        "pressures",
        "recoveries",
        "shot_assists",
        "final_third_passes",
        "final_third_carries",
        "total_actions"
    ]
].copy()


player_display.columns = [
    "Player",
    "Passes",
    "Carries",
    "Shots",
    "xG",
    "Pressures",
    "Recoveries",
    "Shot Assists",
    "Final-Third Passes",
    "Final-Third Carries",
    "Total Actions"
]


add_dataframe(
    ws,
    player_display,
    start_row=4,
    table_name="PlayerMetrics"
)


ws.freeze_panes = "A5"

auto_width(ws)


# xG formatting

for row in range(
    5,
    5 + len(player_display)
):

    ws.cell(
        row=row,
        column=5
    ).number_format = "0.00"


# ============================================================
# SHEET 3 - SQUAD DEPENDENCY
# ============================================================

ws = wb.create_sheet(
    "Squad Dependency"
)

ws.sheet_view.showGridLines = False

add_title(
    ws,
    "BAYER LEVERKUSEN | SQUAD DEPENDENCY",
    "Functional concentration among regularly involved players"
)


dependency_display = dependency[
    [
        "player",
        "attacking_dependency",
        "possession_dependency",
        "defensive_dependency"
    ]
].copy()


dependency_display.columns = [
    "Player",
    "Attacking Dependency",
    "Possession Dependency",
    "Defensive Dependency"
]


dependency_display = dependency_display.sort_values(
    "Attacking Dependency",
    ascending=False
)


add_dataframe(
    ws,
    dependency_display,
    start_row=4,
    table_name="SquadDependency"
)


ws.freeze_panes = "A5"

auto_width(ws)


# Dependency values are already percentage points,
# so display the % symbol without multiplying by 100.

for row in range(
    5,
    5 + len(dependency_display)
):

    for col in range(2, 5):

        ws.cell(
            row=row,
            column=col
        ).number_format = '0.0"%"'


# Highlight top attacking dependency

ws["A5"].fill = PatternFill(
    "solid",
    fgColor="FCE8E8"
)

ws["B5"].fill = PatternFill(
    "solid",
    fgColor="FCE8E8"
)

ws["A5"].font = Font(
    bold=True,
    color=DARK_RED
)

ws["B5"].font = Font(
    bold=True,
    color=DARK_RED
)


# ============================================================
# SHEET 4 - RECRUITMENT PROFILE
# ============================================================

ws = wb.create_sheet(
    "Recruitment Profile"
)

ws.sheet_view.showGridLines = False

add_title(
    ws,
    "BAYER LEVERKUSEN | RECRUITMENT PRIORITY PROFILE",
    "Attacking functions most dependent on Florian Wirtz"
)


profile = pd.DataFrame({
    "Recruitment Attribute": [
        "Final-Third Carrying",
        "Chance Creation",
        "Final-Third Passing",
        "Shot / xG Threat"
    ],

    "Team Output Share": [
        15.58,
        14.70,
        12.90,
        11.65
    ],

    "Recruitment Interpretation": [
        "Prioritize players comfortable progressing possession through carries.",
        "Target players capable of consistently creating shooting opportunities.",
        "Prioritize progressive passing into advanced attacking areas.",
        "Maintain a credible scoring and expected-goals threat."
    ]
})


add_dataframe(
    ws,
    profile,
    start_row=4,
    table_name="RecruitmentProfile"
)


for row in range(
    5,
    5 + len(profile)
):

    ws.cell(
        row=row,
        column=2
    ).number_format = '0.0"%"'

    ws.cell(
        row=row,
        column=3
    ).alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )


ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 70

ws.freeze_panes = "A5"


# ============================================================
# SHEET 5 - VISUAL ANALYSIS
# ============================================================

ws = wb.create_sheet(
    "Visual Analysis"
)

ws.sheet_view.showGridLines = False

add_title(
    ws,
    "BAYER LEVERKUSEN | VISUAL ANALYSIS",
    "Squad dependency, progression and recruitment findings"
)


chart_files = [
    (
        "images/attacking_dependency.png",
        "A4"
    ),
    (
        "images/final_third_progression.png",
        "A29"
    ),
    (
        "images/squad_risk_matrix.png",
        "A54"
    ),
    (
        "images/recruitment_priority_profile.png",
        "A79"
    )
]


for file_path, position in chart_files:

    if os.path.exists(file_path):

        img = Image(file_path)

        # Keep charts large enough to read
        img.width = 900
        img.height = 530

        ws.add_image(
            img,
            position
        )

    else:

        print(
            f"Warning: could not find {file_path}"
        )


for col in range(1, 15):

    ws.column_dimensions[
        get_column_letter(col)
    ].width = 12


# ============================================================
# SHEET 6 - METHODOLOGY
# ============================================================

ws = wb.create_sheet(
    "Methodology"
)

ws.sheet_view.showGridLines = False

add_title(
    ws,
    "BAYER LEVERKUSEN | METHODOLOGY",
    "Metric definitions and analytical limitations"
)


methodology = [
    (
        "Data Source",
        "StatsBomb Open Data covering Bayer Leverkusen's "
        "2023/24 Bundesliga matches."
    ),

    (
        "Final Third",
        "The attacking final third begins at x = 80 on "
        "StatsBomb's 120 × 80 pitch."
    ),

    (
        "Attacking Dependency",
        "Equal-weight average of a player's team share of xG, "
        "shot assists, final-third passes and final-third carries."
    ),

    (
        "Possession Dependency",
        "Equal-weight average of team share of passes and carries."
    ),

    (
        "Defensive Dependency",
        "Equal-weight average of team share of pressures "
        "and ball recoveries."
    ),

    (
        "Main Squad Filter",
        "Players with at least 500 measured actions were included "
        "in the primary dependency analysis."
    ),

    (
        "Interpretation",
        "Dependency represents concentration of measured team "
        "functions. It should not be interpreted as injury risk "
        "or a complete estimate of player replacement value."
    ),

    (
        "Recruitment Profile",
        "The profile identifies the attacking functions most "
        "concentrated in the highest-dependency player. It does "
        "not imply that recruitment should seek an identical player."
    )
]


ws["A4"] = "Method"
ws["B4"] = "Definition"

style_header(ws["A4"])
style_header(ws["B4"])


for row_num, (
    method,
    definition
) in enumerate(
    methodology,
    start=5
):

    ws.cell(
        row=row_num,
        column=1,
        value=method
    )

    ws.cell(
        row=row_num,
        column=2,
        value=definition
    )

    ws.cell(
        row=row_num,
        column=1
    ).font = Font(
        bold=True,
        color=DARK_RED
    )

    ws.cell(
        row=row_num,
        column=2
    ).alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )

    ws.row_dimensions[
        row_num
    ].height = 38


ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 90


# ============================================================
# WORKBOOK PROPERTIES
# ============================================================

wb.properties.title = (
    "Bayer Leverkusen Squad Risk & Recruitment Analysis"
)

wb.properties.subject = (
    "2023/24 Bundesliga Squad Analytics"
)

wb.properties.creator = "Joshua Barnum"

wb.properties.description = (
    "Event-level squad dependency and recruitment analysis "
    "using StatsBomb Open Data."
)


# Make Executive Summary the opening sheet

wb.active = 0


# ============================================================
# SAVE WORKBOOK
# ============================================================

wb.save(
    OUTPUT_FILE
)


print("\n" + "=" * 60)

print(
    "EXCEL REPORT CREATED SUCCESSFULLY"
)

print("=" * 60)

print(
    f"\nSaved as: {OUTPUT_FILE}"
)

print(
    "\nSheets:"
)

print("  1. Executive Summary")
print("  2. Player Metrics")
print("  3. Squad Dependency")
print("  4. Recruitment Profile")
print("  5. Visual Analysis")
print("  6. Methodology")